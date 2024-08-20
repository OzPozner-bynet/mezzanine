""" excel functions """

#my excel functions
#import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
#from openpyxl.utils import get_column_letter
from .models import Student, Course, CourseCategory, StudentCourseMapping
import re

def autofit_column_widths(worksheet):
    """Autofits column widths based on content in the worksheet."""
    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter  # Get the column letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
                    
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2  # Adjust width as needed
        worksheet.column_dimensions[column_letter].width = adjusted_width


def sanitize_sheet_name(name):
    """Removes invalid characters from a sheet name."""
    invalid_chars = r'[\\/*?\[\]:\|<>]'  # Replace with specific invalid chars
    return re.sub(invalid_chars, '_', name)  # Replace with underscore or desired char

def populate_student_sheet(worksheet):
    """ define the logic to populate studemt sheet"""    
    # Student data
    worksheet.cell(row=1, column=1, value="Display_name")
    worksheet.cell(row=1, column=2, value="User")
    row = 2
    for student in Student.objects.all():
        worksheet.cell(row=row, column=1, value=student.display_name)
        worksheet.cell(row=row, column=2, value=str(student.user))


def populate_course_sheet(worksheet, category):
    """ popluate the courses sheets 1 per category"""
    # Header row
    worksheet.cell(row=1, column=1, value="Student")
    column = 2
    desired_category = CourseCategory.objects.get(name=category)
    for course in Course.objects.filter(category=desired_category):
        worksheet.cell(row=1, column=column, value=course.name)
        column += 1

    # Student data
    row = 2
    for student in Student.objects.all():
        worksheet.cell(row=row, column=1, value=student.display_name)
        column = 2
        desired_category = CourseCategory.objects.get(name=category)
        for course in Course.objects.filter(category=desired_category):
            mapping = StudentCourseMapping.objects.filter(student=student, course=course).first()
            if mapping:
                worksheet.cell(row=row, column=column, value=mapping.status)
            column += 1
        row += 1



def populate_dashboard_sheet(worksheet):
    """ fill the dashboard sheet"""
    # Header row
    worksheet.cell(row=1, column=1, value="Student")
    column = 2
    for category in CourseCategory.objects.all():
        worksheet.cell(row=1, column=column, value=category.name)
        column += 1

    # Student data
    row = 2
    for student in Student.objects.all():
        worksheet.cell(row=row, column=1, value=student.display_name)
        column = 2
        for category in CourseCategory.objects.all():
            completed_courses = StudentCourseMapping.objects.filter(
                student=student,
                course__category=category,
                status='completed'
            ).count()
            worksheet.cell(row=row, column=column, value=completed_courses)
            column += 1
        row += 1


def apply_conditional_formatting(worksheet):
    """ aplly cell format accourding to value"""
    # Assuming status values are 'completed', 'planned', and 'in_progress'
    green_fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    red_fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
    gray_fill = PatternFill(start_color='CECECE', end_color='CECECE', fill_type='solid')
    cyan_fill = PatternFill(start_color='F0FFFF', end_color='F0FFFF', fill_type='solid')
    thin_border = Side(border_style='thin', color='000000')
    font_size = Font(size=12)
    border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
  

    row = 1
    for col in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=row, column=col)
        cell.font = Font(name='Arial')
        cell.alignment = Alignment(horizontal='center')
        cell.fill = gray_fill
        cell.border = border
        cell.font = font_size

    for row in range(2, worksheet.max_row + 1):
        col = 1
        cell = worksheet.cell(row=row, column=col)
        cell.font = Font(name='Arial')
        cell.alignment = Alignment(horizontal='center')
        cell.fill = cyan_fill
        cell.border = border
        cell.font = font_size
        for col in range(2, worksheet.max_column + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.border = border
            cell.font = font_size
            cell.alignment = Alignment(horizontal='center')
            if cell.value == 'completed' or cell.value == '1':
                cell.fill = green_fill
                cell.font = Font(name='Arial')
            elif cell.value == 'planned':
                cell.fill = yellow_fill
            elif cell.value == 'in_progress':
                cell.fill = red_fill
